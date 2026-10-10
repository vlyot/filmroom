"""Holdout candidates: possessions for Kai to label blind from video.

Candidates are chosen only by length, completed passes and how the possession ended (never by any engine
output), so labelling them stays independent of the engine. This is a rough possession split; the real
splitter is built in Phase 1.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from filmroom.data.catalog import MatchData
from filmroom.data.positions import BALL_PID
from filmroom.data.textlog import PassComplete, Shot
from filmroom.eval.labels import (
    COL_BETTER,
    COL_CONFIDENCE,
    COL_FINAL,
    COL_PLAYER,
    COL_TP,
    COL_USE,
    LabelRecord,
)

Ending = Literal["shot", "lost_final_third", "lost_midfield"]
QUOTAS: dict[Ending, int] = {"shot": 6, "lost_final_third": 8, "lost_midfield": 6}
MAX_GAP_FRAMES = 300  # ball loose for more than ~3 s ends the possession
EXCLUDE_MARGIN_FRAMES = 500  # stay ~5 s clear of the tuning set's windows
PITCH_X = 110.0


@dataclass(frozen=True)
class Candidate:
    match_id: str
    team: int
    start_frame: int
    end_frame: int
    start_ms: int
    end_ms: int
    passes: int
    ending: Ending


def _team(pids: np.ndarray) -> np.ndarray:
    return (pids > BALL_PID).astype(np.int8)


def _zone(x: float, team: int) -> str:
    progress = x / PITCH_X if team == 0 else 1 - x / PITCH_X
    return "final" if progress > 2 / 3 else "mid" if progress >= 1 / 3 else "own"


def find_candidates(md: MatchData, min_s: float = 6.0, min_passes: int = 3) -> list[Candidate]:
    frames, teams = md.possession_frames, _team(md.possession_pids)
    if len(frames) == 0:
        return []
    breaks = np.flatnonzero((np.diff(teams) != 0) | (np.diff(frames) > MAX_GAP_FRAMES)) + 1
    passes = [e for e in md.events if isinstance(e, PassComplete)]
    shots = [e for e in md.events if isinstance(e, Shot)]
    ball = list(md.track.pids).index(BALL_PID)
    out = []
    for lo, hi in zip(np.r_[0, breaks], np.r_[breaks, len(frames)], strict=True):
        start, end, team = int(frames[lo]), int(frames[hi - 1]), int(teams[lo])
        start_ms, end_ms = md.clock.ms_at(start), md.clock.ms_at(end)
        n_passes = sum(start <= p.start_frame <= end and (p.pid > BALL_PID) == bool(team) for p in passes)
        if end_ms - start_ms < min_s * 1000 or n_passes < min_passes:
            continue
        if any(start <= s.start_frame <= end + 50 and (s.pid > BALL_PID) == bool(team) for s in shots):
            ending: Ending = "shot"
        else:
            idx = min(int(np.searchsorted(md.track.frames, end)), len(md.track.frames) - 1)
            zone = _zone(float(md.track.xy[idx, ball, 0]), team)
            if zone == "own":
                continue
            ending = "lost_final_third" if zone == "final" else "lost_midfield"
        out.append(Candidate(md.info.id, team, start, end, start_ms, end_ms, n_passes, ending))
    return out


def select_candidates(
    candidates: list[Candidate], exclude: list[LabelRecord], quotas: dict[Ending, int] = QUOTAS
) -> list[Candidate]:
    """Fill each ending's quota round-robin across matches, skipping anything near a tuning possession."""

    def clear(c: Candidate) -> bool:
        return not any(
            r.match_id == c.match_id
            and c.start_frame <= r.end_frame + EXCLUDE_MARGIN_FRAMES
            and r.start_frame - EXCLUDE_MARGIN_FRAMES <= c.end_frame
            for r in exclude
        )

    chosen = []
    for ending, quota in quotas.items():
        by_match: dict[str, list[Candidate]] = {}
        for c in sorted(candidates, key=lambda c: (c.match_id, c.start_frame)):
            if c.ending == ending and clear(c):
                by_match.setdefault(c.match_id, []).append(c)
        picked: list[Candidate] = []
        while len(picked) < quota and any(by_match.values()):
            for queue in by_match.values():
                if queue and len(picked) < quota:
                    picked.append(queue.pop(len(queue) // 2))
        chosen.extend(picked)
    return chosen


def _mmss(ms: float) -> str:
    total = max(0, round(ms / 1000))
    return f"{total // 60:02d}:{total % 60:02d}"


INSTRUCTIONS = [
    "Holdout possessions — label blind from video",
    "",
    "These possessions were picked only by length, passes and how they ended. No engine output was used.",
    "Choose 10 rows (3 shot, 4 lost in the final third, 3 lost in midfield) by putting Y in 'Use? (Y)'.",
    "For each chosen row, watch the video from 'Video start' to 'Video end' in the match's recording.mkv,",
    "then fill the green columns. Do not look at any FilmRoom output first.",
    "Turning point: the video time (mm:ss) of the decision where the move went wrong. Leave blank if none.",
    "Player: the small number shown above the player in the video, e.g. #151.",
    "Better option: name the player as #NNN and describe the option.",
]

COLUMNS = [
    "ID", COL_USE, "Match", "Mode", "Team", "Game start", "Game end", "Video start", "Video end",
    "Length (s)", "Completed passes", "Ended (auto)", COL_TP, COL_PLAYER, "What they did", COL_BETTER,
    "Why it was better", COL_FINAL, COL_CONFIDENCE, "Notes", "start_frame", "end_frame",
]  # fmt: skip
LABEL_COLUMNS = {COL_USE, COL_TP, COL_PLAYER, "What they did", COL_BETTER, "Why it was better", COL_FINAL,
                 COL_CONFIDENCE, "Notes"}  # fmt: skip


def write_holdout_sheet(
    candidates: list[Candidate], modes: dict[str, str], offsets: dict[str, float], path: Path
) -> None:
    if path.exists():
        raise FileExistsError(f"{path} already exists; it may hold labels, so it is never overwritten")
    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = "Instructions"
    for line in INSTRUCTIONS:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=13)
    intro.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Possessions")
    ws.append(COLUMNS)
    green = PatternFill("solid", fgColor="E2EFDA")
    for i, c in enumerate(candidates, start=1):
        offset_ms = offsets[c.match_id] * 1000
        side = f"Team {c.team} (attacks → x={'110' if c.team == 0 else '0'})"
        ws.append([
            f"H{i:02d}", None, c.match_id, modes.get(c.match_id, "unknown"), side,
            _mmss(c.start_ms), _mmss(c.end_ms), _mmss(c.start_ms + offset_ms), _mmss(c.end_ms + offset_ms),
            round((c.end_ms - c.start_ms) / 1000), c.passes, c.ending.replace("_", " "),
            None, None, None, None, None, None, None, None, c.start_frame, c.end_frame,
        ])  # fmt: skip
    for j, name in enumerate(COLUMNS, start=1):
        ws.cell(1, j).font = Font(bold=True)
        ws.cell(1, j).alignment = Alignment(wrap_text=True)
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 14
        if name in LABEL_COLUMNS:
            for row in range(1, len(candidates) + 2):
                ws.cell(row, j).fill = green
    ws.freeze_panes = "C2"

    offsets_ws = wb.create_sheet("Offsets")
    offsets_ws.append(["Match", "Mode", "Kick-off offset in video (s)"])
    for match_id, offset in offsets.items():
        offsets_ws.append([match_id, modes.get(match_id, "unknown"), offset])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
