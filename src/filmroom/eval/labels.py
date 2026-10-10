"""Load the labelled reference sets (tuning and holdout spreadsheets) into typed records."""

import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import openpyxl

from filmroom.data.frame_times import FrameClock

COL_TP = "Turning point (video mm:ss)"
COL_PLAYER = "Player"
COL_BETTER = "Better option"
COL_FINAL = "Turning point = final action? (Y/N)"
COL_CONFIDENCE = "Confidence (1–5)"
COL_USE = "Use? (Y)"


@dataclass(frozen=True)
class LabelRecord:
    id: str
    set: Literal["tuning", "holdout"]
    match_id: str
    start_frame: int
    end_frame: int
    tp_video_s: int | None  # None = labeller saw no clear turning point
    tp_frame: int | None
    tp_pid: int | None
    better_pid: int | None
    final_action: bool | None
    confidence: int | None


def mmss_to_s(value: str) -> int:
    minutes, seconds = value.strip().split(":")
    return int(minutes) * 60 + int(seconds)


def _pid(text: str | None) -> int | None:
    match = re.search(r"#(\d+)", text or "")
    return int(match.group(1)) if match else None


def load_offsets(path: Path) -> dict[str, float]:
    ws = openpyxl.load_workbook(path, read_only=True)["Offsets"]
    return {row[0]: float(row[2]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}


def load_labels(
    path: Path,
    set_name: Literal["tuning", "holdout"],
    clock_for: Callable[[str], FrameClock] | None = None,
) -> list[LabelRecord]:
    """Read the Possessions sheet. Turning points are converted from video time to a game frame when a
    clock is available (video time = game time + the match's offset on the Offsets sheet)."""
    wb = openpyxl.load_workbook(path, read_only=True)
    offsets = load_offsets(path)
    rows = wb["Possessions"].iter_rows(values_only=True)
    header = list(next(rows))
    col = {name: i for i, name in enumerate(header)}
    records = []
    for row in rows:
        rid = row[col["ID"]]
        if not rid or rid == "EXAMPLE":
            continue
        if COL_USE in col and str(row[col[COL_USE]] or "").strip().upper() != "Y":
            continue
        match_id = row[col["Match"]]
        tp_text = row[col[COL_TP]]
        tp_video_s = mmss_to_s(tp_text) if tp_text else None
        tp_frame = None
        if tp_video_s is not None and clock_for is not None:
            tp_frame = clock_for(match_id).frame_at(round((tp_video_s - offsets[match_id]) * 1000))
        final = row[col[COL_FINAL]]
        confidence = row[col[COL_CONFIDENCE]]
        records.append(
            LabelRecord(
                id=rid,
                set=set_name,
                match_id=match_id,
                start_frame=int(row[col["start_frame"]]),
                end_frame=int(row[col["end_frame"]]),
                tp_video_s=tp_video_s,
                tp_frame=tp_frame,
                tp_pid=_pid(row[col[COL_PLAYER]]),
                better_pid=_pid(row[col[COL_BETTER]]),
                final_action=None if final is None else str(final).strip().upper() == "Y",
                confidence=int(confidence) if confidence is not None else None,
            )
        )
    return records
