import openpyxl
import pytest

from filmroom.eval.holdout import Candidate, select_candidates, write_holdout_sheet
from filmroom.eval.labels import COL_PLAYER, COL_TP, COL_USE, LabelRecord, load_labels, mmss_to_s


def _cand(match: str, start: int, ending: str) -> Candidate:
    return Candidate(match, 0, start, start + 700, start * 10, start * 10 + 7000, 4, ending)


def _label(match: str, start: int, end: int) -> LabelRecord:
    return LabelRecord("P01", "tuning", match, start, end, None, None, None, None, None, None)


def test_mmss_to_seconds():
    assert mmss_to_s("03:20") == 200
    assert mmss_to_s("00:07") == 7


def test_selection_excludes_windows_near_tuning_possessions():
    near = _cand("M1", 10_400, "shot")  # starts 400 frames after a tuning window ends
    far = _cand("M1", 20_000, "shot")
    chosen = select_candidates([near, far], exclude=[_label("M1", 9_000, 10_000)], quotas={"shot": 2})
    assert chosen == [far]


def test_selection_fills_quotas_round_robin_across_matches():
    cands = [
        _cand(m, s, e) for m in ("M1", "M2") for s in (1_000, 5_000, 9_000) for e in ("shot", "lost_midfield")
    ]
    chosen = select_candidates(cands, exclude=[], quotas={"shot": 2, "lost_midfield": 3})
    assert [c.ending for c in chosen].count("shot") == 2
    assert [c.ending for c in chosen].count("lost_midfield") == 3
    assert {c.match_id for c in chosen if c.ending == "shot"} == {"M1", "M2"}


def test_holdout_sheet_round_trips_through_load_labels(tmp_path):
    path = tmp_path / "holdout.xlsx"
    cands = [_cand("M1", 1_000, "shot"), _cand("M1", 5_000, "lost_midfield")]
    write_holdout_sheet(cands, {"M1": "human vs human"}, {"M1": 4.4}, path)

    wb = openpyxl.load_workbook(path)
    ws = wb["Possessions"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index(COL_USE) + 1, "Y")
    ws.cell(2, header.index(COL_TP) + 1, "01:05")
    ws.cell(2, header.index(COL_PLAYER) + 1, "#151")
    wb.save(path)

    [record] = load_labels(path, "holdout")
    assert (record.id, record.set, record.tp_video_s, record.tp_pid) == ("H01", "holdout", 65, 151)
    assert (record.start_frame, record.end_frame) == (1_000, 1_700)


def test_holdout_sheet_is_never_overwritten(tmp_path):
    path = tmp_path / "holdout.xlsx"
    path.write_bytes(b"existing labels")
    with pytest.raises(FileExistsError):
        write_holdout_sheet([], {}, {}, path)
    assert path.read_bytes() == b"existing labels"
